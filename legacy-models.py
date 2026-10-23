import os
import requests
import pandas as pd
from datetime import datetime, timezone
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

DATABRICKS_HOST = os.getenv("DATABRICKS_HOST")
DATABRICKS_TOKEN = os.getenv("DATABRICKS_TOKEN")

OUTPUT_FILE = "databricks_legacy_model_permissions.xlsx"

if not DATABRICKS_HOST:
    raise ValueError("DATABRICKS_HOST environment variable is not set")

if not DATABRICKS_TOKEN:
    raise ValueError("DATABRICKS_TOKEN environment variable is not set")

DATABRICKS_HOST = DATABRICKS_HOST.rstrip("/")

HEADERS = {
    "Authorization": f"Bearer {DATABRICKS_TOKEN}",
    "Content-Type": "application/json"
}


# ============================================================
# HELPER
# ============================================================

def api_get(url, params=None):
    response = requests.get(
        url,
        headers=HEADERS,
        params=params,
        timeout=60
    )

    if response.status_code != 200:
        raise RuntimeError(
            f"API call failed\n"
            f"URL: {url}\n"
            f"Status: {response.status_code}\n"
            f"Response: {response.text}"
        )

    return response.json()


def epoch_to_datetime(epoch_ms):
    if not epoch_ms:
        return ""

    return datetime.fromtimestamp(
        epoch_ms / 1000,
        tz=timezone.utc
    ).strftime("%Y-%m-%d %H:%M:%S UTC")


# ============================================================
# GET ALL LEGACY REGISTERED MODELS
# ============================================================

def get_all_models():

    url = (
        f"{DATABRICKS_HOST}"
        "/api/2.0/mlflow/registered-models/list"
    )

    models = []
    page_token = None

    while True:

        params = {
            "max_results": 1000
        }

        if page_token:
            params["page_token"] = page_token

        data = api_get(url, params)

        models.extend(
            data.get("registered_models", [])
        )

        page_token = data.get("next_page_token")

        if not page_token:
            break

    return models


# ============================================================
# GET DATABRICKS-SPECIFIC MODEL INFORMATION
# ============================================================

def get_model_details(model_name):

    url = (
        f"{DATABRICKS_HOST}"
        "/api/2.0/mlflow/databricks/"
        "registered-models/get"
    )

    data = api_get(
        url,
        params={"name": model_name}
    )

    return data.get(
        "registered_model_databricks",
        {}
    )


# ============================================================
# GET MODEL PERMISSIONS
# ============================================================

def get_model_permissions(model_id):

    url = (
        f"{DATABRICKS_HOST}"
        f"/api/2.0/permissions/"
        f"registered-models/{model_id}"
    )

    return api_get(url)


# ============================================================
# EXTRACT PRINCIPAL INFORMATION
# ============================================================

def extract_principal(acl):

    if acl.get("user_name"):
        return "User", acl["user_name"]

    if acl.get("group_name"):
        return "Group", acl["group_name"]

    if acl.get("service_principal_name"):
        return (
            "Service Principal",
            acl["service_principal_name"]
        )

    if acl.get("display_name"):
        return "Unknown", acl["display_name"]

    return "Unknown", ""


# ============================================================
# MAIN COLLECTION
# ============================================================

def collect_data():

    print("Getting legacy registered models...")

    models = get_all_models()

    print(f"Found {len(models)} legacy models")

    model_summary = []
    permissions_rows = []

    for index, model in enumerate(models, start=1):

        model_name = model.get("name")

        print(
            f"[{index}/{len(models)}] "
            f"Processing: {model_name}"
        )

        try:

            details = get_model_details(model_name)

            model_id = details.get("id", "")
            creator = details.get("user_id", "")

            creation_timestamp = details.get(
                "creation_timestamp"
            )

            last_updated_timestamp = details.get(
                "last_updated_timestamp"
            )

            permission_level = details.get(
                "permission_level",
                ""
            )

            # ------------------------------------------------
            # GET ACL
            # ------------------------------------------------

            permission_data = {}

            if model_id:
                permission_data = get_model_permissions(
                    model_id
                )

            acl_list = permission_data.get(
                "access_control_list",
                []
            )

            # ------------------------------------------------
            # MODEL SUMMARY
            # ------------------------------------------------

            model_summary.append({
                "Model Name": model_name,
                "Model ID": model_id,
                "Owner / Creator": creator,
                "Created Time": epoch_to_datetime(
                    creation_timestamp
                ),
                "Last Updated": epoch_to_datetime(
                    last_updated_timestamp
                ),
                "My Permission Level": permission_level,
                "Permission Entries": len(acl_list)
            })

            # ------------------------------------------------
            # PERMISSIONS
            # ------------------------------------------------

            for acl in acl_list:

                principal_type, principal = (
                    extract_principal(acl)
                )

                permissions = acl.get(
                    "all_permissions",
                    []
                )

                for permission in permissions:

                    permissions_rows.append({
                        "Model Name": model_name,
                        "Model ID": model_id,
                        "Principal Type": principal_type,
                        "Principal": principal,
                        "Permission": permission.get(
                            "permission_level",
                            ""
                        ),
                        "Inherited": permission.get(
                            "inherited",
                            False
                        ),
                        "Inherited From": ", ".join(
                            permission.get(
                                "inherited_from_object",
                                []
                            )
                        )
                    })

        except Exception as e:

            print(
                f"ERROR processing {model_name}: {e}"
            )

            model_summary.append({
                "Model Name": model_name,
                "Model ID": "",
                "Owner / Creator": "",
                "Created Time": "",
                "Last Updated": "",
                "My Permission Level": "",
                "Permission Entries": "",
                "Error": str(e)
            })

    return model_summary, permissions_rows


# ============================================================
# WRITE EXCEL
# ============================================================

def create_excel(model_summary, permissions_rows):

    df_summary = pd.DataFrame(model_summary)

    df_permissions = pd.DataFrame(
        permissions_rows
    )

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        df_summary.to_excel(
            writer,
            sheet_name="Model Summary",
            index=False
        )

        df_permissions.to_excel(
            writer,
            sheet_name="Permissions",
            index=False
        )

    # --------------------------------------------------------
    # FORMAT EXCEL
    # --------------------------------------------------------

    workbook = load_workbook(
        OUTPUT_FILE
    )

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        # Header formatting
        for cell in worksheet[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # Auto-size columns
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                60
            )

    workbook.save(
        OUTPUT_FILE
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    summary, permissions = collect_data()

    create_excel(
        summary,
        permissions
    )

    print()
    print("=" * 60)
    print("COMPLETED")
    print("=" * 60)
    print(f"Models found      : {len(summary)}")
    print(
        f"Permission records: {len(permissions)}"
    )
    print(f"Excel file        : {OUTPUT_FILE}")
