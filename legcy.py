# ============================================================
# DATABRICKS LEGACY MODEL REGISTRY
# OWNER + CURRENT PERMISSIONS AUDIT
#
# Runs directly inside a Databricks Notebook
# No PAT or DATABRICKS_HOST environment variable required
# ============================================================

import requests
import pandas as pd

from datetime import datetime, timezone

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# 1. GET DATABRICKS NOTEBOOK AUTHENTICATION
# ============================================================

print("=" * 70)
print("DATABRICKS LEGACY MODEL PERMISSION AUDIT")
print("=" * 70)

try:

    DATABRICKS_HOST = (
        dbutils.notebook
        .entry_point
        .getDbutils()
        .notebook()
        .getContext()
        .apiUrl()
        .get()
    )

    DATABRICKS_TOKEN = (
        dbutils.notebook
        .entry_point
        .getDbutils()
        .notebook()
        .getContext()
        .apiToken()
        .get()
    )

except Exception as e:

    raise RuntimeError(
        "Unable to obtain Databricks notebook authentication. "
        f"Error: {e}"
    )


DATABRICKS_HOST = DATABRICKS_HOST.rstrip("/")


HEADERS = {
    "Authorization": f"Bearer {DATABRICKS_TOKEN}",
    "Content-Type": "application/json"
}


print(f"Databricks Host : {DATABRICKS_HOST}")
print("Authentication   : SUCCESS")
print()


# ============================================================
# 2. OUTPUT FILE
# ============================================================

OUTPUT_FILE = (
    "/dbfs/FileStore/"
    "databricks_legacy_model_permissions.xlsx"
)


# ============================================================
# 3. GENERIC API GET FUNCTION
# ============================================================

def api_get(url, params=None):

    response = requests.get(
        url,
        headers=HEADERS,
        params=params,
        timeout=60
    )

    if response.status_code != 200:

        raise RuntimeError(
            "\nDatabricks API request failed\n"
            f"URL       : {url}\n"
            f"Status    : {response.status_code}\n"
            f"Response  : {response.text}"
        )

    return response.json()


# ============================================================
# 4. CONVERT EPOCH MILLISECONDS TO UTC
# ============================================================

def epoch_to_datetime(epoch_ms):

    if not epoch_ms:
        return ""

    try:

        return datetime.fromtimestamp(
            epoch_ms / 1000,
            tz=timezone.utc
        ).strftime(
            "%Y-%m-%d %H:%M:%S UTC"
        )

    except Exception:

        return str(epoch_ms)


# ============================================================
# 5. GET ALL LEGACY REGISTERED MODELS
# ============================================================

def get_all_legacy_models():

    print("Getting legacy Workspace Model Registry models...")

    url = (
        f"{DATABRICKS_HOST}"
        "/api/2.0/mlflow/registered-models/list"
    )

    models = []

    page_token = None

    while True:

        params = {
            "max_results": 1000
        }

        if page_token:
            params["page_token"] = page_token

        data = api_get(
            url,
            params
        )

        registered_models = data.get(
            "registered_models",
            []
        )

        models.extend(
            registered_models
        )

        page_token = data.get(
            "next_page_token"
        )

        if not page_token:
            break

    print(
        f"Legacy models found: {len(models)}"
    )

    print()

    return models


# ============================================================
# 6. GET LEGACY MODEL DETAILS
# ============================================================

def get_model_details(model_name):

    url = (
        f"{DATABRICKS_HOST}"
        "/api/2.0/mlflow/databricks/"
        "registered-models/get"
    )

    return api_get(
        url,
        params={
            "name": model_name
        }
    )


# ============================================================
# 7. GET LEGACY MODEL PERMISSIONS
# ============================================================

def get_model_permissions(model_id):

    url = (
        f"{DATABRICKS_HOST}"
        "/api/2.0/permissions/"
        f"registered-models/{model_id}"
    )

    return api_get(url)


# ============================================================
# 8. IDENTIFY PRINCIPAL
# ============================================================

def extract_principal(acl):

    # User
    if acl.get("user_name"):

        return (
            "User",
            acl.get("user_name")
        )

    # Group
    if acl.get("group_name"):

        return (
            "Group",
            acl.get("group_name")
        )

    # Service Principal
    if acl.get("service_principal_name"):

        return (
            "Service Principal",
            acl.get("service_principal_name")
        )

    # Some APIs can return display_name
    if acl.get("display_name"):

        return (
            "Unknown",
            acl.get("display_name")
        )

    return (
        "Unknown",
        ""
    )


# ============================================================
# 9. COLLECT MODEL INFORMATION
# ============================================================

def collect_model_data():

    models = get_all_legacy_models()

    model_summary = []

    permissions_rows = []

    errors = []

    total_models = len(models)

    print("=" * 70)
    print("PROCESSING MODELS")
    print("=" * 70)

    for index, model in enumerate(
        models,
        start=1
    ):

        model_name = model.get(
            "name",
            ""
        )

        print(
            f"[{index}/{total_models}] "
            f"{model_name}"
        )

        try:

            # ------------------------------------------------
            # GET MODEL DETAILS
            # ------------------------------------------------

            details_response = (
                get_model_details(
                    model_name
                )
            )

            details = (
                details_response.get(
                    "registered_model_databricks",
                    {}
                )
            )

            model_id = details.get(
                "id",
                ""
            )

            owner = details.get(
                "user_id",
                ""
            )

            created_timestamp = (
                details.get(
                    "creation_timestamp"
                )
            )

            updated_timestamp = (
                details.get(
                    "last_updated_timestamp"
                )
            )

            permission_level = (
                details.get(
                    "permission_level",
                    ""
                )
            )

            # ------------------------------------------------
            # GET CURRENT PERMISSIONS
            # ------------------------------------------------

            permission_data = {}

            if model_id:

                permission_data = (
                    get_model_permissions(
                        model_id
                    )
                )

            acl_list = (
                permission_data.get(
                    "access_control_list",
                    []
                )
            )

            # ------------------------------------------------
            # MODEL SUMMARY
            # ------------------------------------------------

            model_summary.append({

                "Registry Type":
                    "LEGACY",

                "Model Name":
                    model_name,

                "Model ID":
                    model_id,

                "Owner / Creator":
                    owner,

                "Created Time":
                    epoch_to_datetime(
                        created_timestamp
                    ),

                "Last Updated":
                    epoch_to_datetime(
                        updated_timestamp
                    ),

                "My Permission Level":
                    permission_level,

                "Permission Entries":
                    len(acl_list)

            })

            # ------------------------------------------------
            # CURRENT PERMISSIONS
            # ------------------------------------------------

            for acl in acl_list:

                (
                    principal_type,
                    principal
                ) = extract_principal(
                    acl
                )

                all_permissions = (
                    acl.get(
                        "all_permissions",
                        []
                    )
                )

                # --------------------------------------------
                # PERMISSIONS FOR PRINCIPAL
                # --------------------------------------------

                for permission in all_permissions:

                    inherited_from = (
                        permission.get(
                            "inherited_from_object",
                            []
                        )
                    )

                    if isinstance(
                        inherited_from,
                        list
                    ):

                        inherited_from = (
                            ", ".join(
                                inherited_from
                            )
                        )

                    permissions_rows.append({

                        "Registry Type":
                            "LEGACY",

                        "Model Name":
                            model_name,

                        "Model ID":
                            model_id,

                        "Principal Type":
                            principal_type,

                        "Principal":
                            principal,

                        "Permission":
                            permission.get(
                                "permission_level",
                                ""
                            ),

                        "Inherited":
                            permission.get(
                                "inherited",
                                False
                            ),

                        "Inherited From":
                            inherited_from

                    })

        except Exception as e:

            error_message = str(e)

            print(
                f"   ERROR: {error_message}"
            )

            errors.append({

                "Registry Type":
                    "LEGACY",

                "Model Name":
                    model_name,

                "Error":
                    error_message

            })

            # Add model to summary even if permission
            # retrieval failed
            model_summary.append({

                "Registry Type":
                    "LEGACY",

                "Model Name":
                    model_name,

                "Model ID":
                    "",

                "Owner / Creator":
                    "",

                "Created Time":
                    "",

                "Last Updated":
                    "",

                "My Permission Level":
                    "",

                "Permission Entries":
                    "",

                "Error":
                    error_message

            })

    return (
        model_summary,
        permissions_rows,
        errors
    )


# ============================================================
# 10. CREATE EXCEL REPORT
# ============================================================

def create_excel_report(
    model_summary,
    permissions_rows,
    errors
):

    print()
    print("=" * 70)
    print("CREATING EXCEL REPORT")
    print("=" * 70)

    # --------------------------------------------------------
    # DATAFRAMES
    # --------------------------------------------------------

    df_models = pd.DataFrame(
        model_summary
    )

    df_permissions = pd.DataFrame(
        permissions_rows
    )

    df_errors = pd.DataFrame(
        errors
    )

    # --------------------------------------------------------
    # PERMISSION SUMMARY
    # --------------------------------------------------------

    permission_summary = []

    if not df_permissions.empty:

        grouped = (
            df_permissions
            .groupby(
                [
                    "Model Name",
                    "Model ID"
                ]
            )
        )

        for (
            model_key,
            group
        ) in grouped:

            model_name = model_key[0]
            model_id = model_key[1]

            users = group[
                group["Principal Type"]
                == "User"
            ]["Principal"].nunique()

            groups = group[
                group["Principal Type"]
                == "Group"
            ]["Principal"].nunique()

            service_principals = group[
                group["Principal Type"]
                == "Service Principal"
            ]["Principal"].nunique()

            permission_summary.append({

                "Model Name":
                    model_name,

                "Model ID":
                    model_id,

                "Users":
                    users,

                "Groups":
                    groups,

                "Service Principals":
                    service_principals,

                "Total Permission Records":
                    len(group)

            })

    df_permission_summary = pd.DataFrame(
        permission_summary
    )

    # --------------------------------------------------------
    # WRITE EXCEL
    # --------------------------------------------------------

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        df_models.to_excel(
            writer,
            sheet_name="Model Summary",
            index=False
        )

        df_permissions.to_excel(
            writer,
            sheet_name="Current Permissions",
            index=False
        )

        df_permission_summary.to_excel(
            writer,
            sheet_name="Permission Summary",
            index=False
        )

        df_errors.to_excel(
            writer,
            sheet_name="Errors",
            index=False
        )

    # --------------------------------------------------------
    # FORMAT EXCEL
    # --------------------------------------------------------

    workbook = load_workbook(
        OUTPUT_FILE
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for worksheet in workbook.worksheets:

        # Freeze first row
        worksheet.freeze_panes = "A2"

        # Filter
        if worksheet.max_row > 1:

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

        # Header formatting
        for cell in worksheet[1]:

            cell.font = header_font

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Auto-size columns
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                get_column_letter(
                    column_cells[0].column
                )
            )

            for cell in column_cells:

                if cell.value is not None:

                    value_length = len(
                        str(cell.value)
                    )

                    if value_length > max_length:
                        max_length = value_length

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                60
            )

    workbook.save(
        OUTPUT_FILE
    )

    print()
    print(
        f"Excel report created successfully:"
    )

    print(
        OUTPUT_FILE
    )


# ============================================================
# 11. RUN AUDIT
# ============================================================

(
    model_summary,
    permissions_rows,
    errors
) = collect_model_data()


# ============================================================
# 12. CREATE EXCEL
# ============================================================

create_excel_report(
    model_summary,
    permissions_rows,
    errors
)


# ============================================================
# 13. FINAL SUMMARY
# ============================================================

print()
print("=" * 70)
print("AUDIT COMPLETED")
print("=" * 70)

print(
    f"Total Models           : "
    f"{len(model_summary)}"
)

print(
    f"Permission Records     : "
    f"{len(permissions_rows)}"
)

print(
    f"Errors                 : "
    f"{len(errors)}"
)

print()
print(
    f"Excel File             : "
    f"{OUTPUT_FILE}"
)

print("=" * 70)
